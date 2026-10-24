"""Chargement d'un jeu de travail depuis la suite Atelier ou un fichier joint.

- Atelier Viz : profil du jeu (`GET /datasets/{id}`, types des colonnes) puis lignes par pages de
  1 000 (`GET /datasets/{id}/rows`), via `DataViaClient` (CLAUDE.md : seul point d'entrée).
- Atelier Tab : instantané DataFlow lu par `TabSnapshots` (valeurs calculées des formules, lignes
  de total exclues), mis au format `atelier-table/1` comme pour les workflows.
- Fichier joint à la conversation (`env.conversation_files()`) : CSV / TSV / TXT délimité,
  Excel `.xlsx` / `.xlsm` (valeurs calculées), JSON (tableau d'objets). Types inférés.
"""

from __future__ import annotations

import csv
import io
import json
from collections.abc import Iterable
from typing import Any

from app.analysis.frame import ColumnType, Frame, frame_from_rows
from app.tools.data_via import DataViaClient
from app.tools.files import ConversationFiles
from app.tools.tab import TabSnapshots, envelope_table, sorted_records

PAGE = 1000
MAX_CELLS = 3_000_000


class LoadError(ValueError):
    """Source illisible ou trop volumineuse : message clair au modèle."""


def _cap_rows(max_rows: int, width: int) -> int:
    return max(1, min(max_rows, MAX_CELLS // max(1, width)))


async def from_viz(client: DataViaClient, dataset_id: str, name: str, max_rows: int) -> Frame:
    profile = await client.get_dataset(dataset_id)
    types: dict[str, ColumnType] = {c.name: c.type for c in profile.columns}
    headers: list[str] | None = None
    rows: list[list[Any]] = []
    total = profile.rows
    limit = _cap_rows(max_rows, len(profile.columns))
    offset = 0
    while offset < min(total, limit):
        page = await client.get_rows(dataset_id, offset=offset, limit=min(PAGE, limit - offset))
        if headers is None:
            headers = page.columns or [c.name for c in profile.columns]
        total = page.total
        rows += [[r.get(h) for h in headers] for r in page.rows[: limit - len(rows)]]
        if not page.rows:
            break
        offset += len(page.rows)
    headers = headers or [c.name for c in profile.columns]
    frame = frame_from_rows(
        name,
        profile.name,
        f"Atelier Viz · jeu « {profile.name} » ({profile.id})",
        headers,
        rows,
        types=[types.get(h, "string") for h in headers],
    )
    frame.source_rows = total
    if len(rows) < total:
        frame.notes.append(f"Chargé : {len(rows)} ligne(s) sur {total} (limite de chargement).")
    return frame


_TAB_TYPES: dict[str, ColumnType] = {
    "string": "string", "number": "number", "date": "date", "datetime": "date", "boolean": "boolean",
}  # fmt: skip


async def from_tab(snapshots: TabSnapshots, table_id: str, name: str, max_rows: int) -> Frame:
    env, info = await snapshots.table(table_id)
    columns = env.project.columns_schema
    records = sorted_records(env.project.records)
    limit = _cap_rows(max_rows, len(columns))
    window = records[:limit]
    table = envelope_table(
        env, info, columns, window, total=len(records), truncated=len(window) < len(records)
    )
    cols = [c for c in table["columns"] if c["key"] not in ("_ligne", "_id")]
    rows = [[r.get(c["key"]) for c in cols] for r in table["rows"]]
    frame = frame_from_rows(
        name,
        info.name,
        f"Atelier Tab · table « {info.name} » (projet {info.project}, {info.id})",
        [c["label"] for c in cols],
        rows,
        types=[_TAB_TYPES.get(c["type"], "string") for c in cols],
    )
    frame.source_rows = len(records)
    if len(window) < len(records):
        frame.notes.append(f"Chargé : {len(window)} ligne(s) sur {len(records)}.")
    if env.project.summary_rows:
        frame.notes.append("Lignes de total de la table exclues (elles fausseraient les calculs).")
    return frame


async def from_file(
    files: ConversationFiles, ref: str, name: str, max_rows: int, sheet: str | None = None
) -> Frame:
    info, data = await files.load(ref)
    ext = info.name.rsplit(".", 1)[-1].lower() if "." in info.name else ""
    keep = max_rows + _HEAD  # lignes gardées en mémoire (les suivantes sont seulement comptées)
    if ext in ("csv", "tsv", "txt"):
        headers, rows, notes, total = _delimited(data, keep)
    elif ext in ("xlsx", "xlsm"):
        headers, rows, notes, total = _excel(data, sheet, keep)
    elif ext == "json":
        headers, rows, notes = _json(data)
        total = len(rows)
    elif ext in ("pdf", "xls", "ofx", "qif", "ods"):
        raise LoadError(
            f"« {info.name} » ne se lit pas directement : convertis-le d'abord en table (Atelier "
            "Convert, puis envoi vers Atelier Viz ou Tab) et charge ensuite ce jeu ; ou demande à "
            "l'utilisateur un export CSV ou .xlsx."
        )
    else:
        raise LoadError(f"Format de « {info.name} » non pris en charge (CSV, TSV, .xlsx, JSON).")
    if not headers:
        raise LoadError(f"Aucun tableau lisible dans « {info.name} ».")
    limit = _cap_rows(max_rows, len(headers))
    frame = frame_from_rows(
        name, info.name, f"Fichier joint « {info.name} »", headers, rows[:limit]
    )
    frame.notes = notes + frame.notes
    frame.source_rows = total
    if total > frame.rows:
        frame.notes.append(
            f"Chargé : {frame.rows} ligne(s) sur environ {total} (limite de chargement)."
        )
    return frame


def _strip(rows: list[list[Any]]) -> list[list[Any]]:
    """Retire lignes vides et colonnes vides en fin de tableau."""

    def empty(v: Any) -> bool:
        return v is None or (isinstance(v, str) and not v.strip())

    rows = [r for r in rows if not all(empty(v) for v in r)]
    width = max((len(r) for r in rows), default=0)
    while width > 0 and all(len(r) < width or empty(r[width - 1]) for r in rows):
        width -= 1
    return [list(r[:width]) + [None] * (width - len(r[:width])) for r in rows]


def _header_index(rows: list[list[Any]]) -> int:
    """Ligne d'en-tête : la première des 20 premières assez remplie (titres au-dessus ignorés)."""
    head = rows[:20]
    counts = [sum(1 for v in r if v not in (None, "")) for r in head]
    best = max(counts, default=0)
    for i, c in enumerate(counts):
        if c >= max(1, 0.6 * best):
            return i
    return 0


def _split(rows: list[list[Any]]) -> tuple[list[str], list[list[Any]], list[str]]:
    rows = _strip(rows)
    if not rows:
        return [], [], []
    h = _header_index(rows)
    notes = [f"{h} ligne(s) de titre ignorée(s) au-dessus de l'en-tête."] if h else []
    headers = ["" if v is None else str(v).strip() for v in rows[h]]
    return headers, rows[h + 1 :], notes


_HEAD = 25  # lignes de titre possibles au-dessus de l'en-tête


def _bounded(rows: Iterable[list[Any]], keep: int) -> tuple[list[list[Any]], int]:
    """Garde les `keep` premières lignes, compte les autres sans les garder en mémoire."""
    kept: list[list[Any]] = []
    total = 0
    for row in rows:
        total += 1
        if len(kept) < keep:
            kept.append(row)
    return kept, total


def _delimited(data: bytes, keep: int) -> tuple[list[str], list[list[Any]], list[str], int]:
    from app.attachments.processing import decode_text

    text = decode_text(data)
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:20_000], delimiters=",;\t|"
        )
    except csv.Error:
        dialect = csv.excel
    rows, total = _bounded((list(r) for r in csv.reader(io.StringIO(text), dialect)), keep)
    headers, body, notes = _split(rows)
    return headers, body, notes, max(len(body), total - (len(rows) - len(body)))


def _excel(
    data: bytes, sheet: str | None, keep: int
) -> tuple[list[str], list[list[Any]], list[str], int]:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise LoadError(f"Classeur Excel illisible : {exc}") from exc
    try:
        names = workbook.sheetnames
        if sheet is not None:
            match = [n for n in names if n.strip().lower() == sheet.strip().lower()]
            if not match:
                raise LoadError(f"Feuille « {sheet} » absente. Feuilles : {', '.join(names)}.")
            ws = workbook[match[0]]
        else:
            ws = workbook[names[0]]
        rows, total = _bounded((list(r) for r in ws.iter_rows(values_only=True)), keep)
    finally:
        workbook.close()
    headers, body, notes = _split(rows)
    if sheet is None and len(names) > 1:
        notes.insert(0, f"Feuille « {names[0]} » lue ; autres feuilles : {', '.join(names[1:])}.")
    return headers, body, notes, max(len(body), total - (len(rows) - len(body)))


def _json(data: bytes) -> tuple[list[str], list[list[Any]], list[str]]:
    try:
        payload = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise LoadError(f"JSON illisible : {exc}") from exc
    if isinstance(payload, dict):  # {"rows": [...]} ou table atelier-table/1
        payload = payload.get("rows") or payload.get("data") or []
    if not isinstance(payload, list) or not all(isinstance(r, dict) for r in payload):
        raise LoadError("JSON attendu : un tableau d'objets (une ligne par objet).")
    headers = list(dict.fromkeys(k for r in payload for k in r))
    rows = [[_scalar(r.get(h)) for h in headers] for r in payload]
    return headers, rows, []


def _scalar(value: Any) -> Any:
    return json.dumps(value, ensure_ascii=False) if isinstance(value, dict | list) else value
