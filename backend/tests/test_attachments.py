"""Étape 15 : pièces jointes dans les conversations."""

from __future__ import annotations

import base64
import io
import json
from collections.abc import AsyncIterator, Iterator
from datetime import UTC, datetime, timedelta
from pathlib import Path

import docx
import openpyxl
import pytest
from fastapi.testclient import TestClient
from PIL import Image
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from app.agents import AgentCatalog
from app.api.deps import (
    get_attachment_repository,
    get_catalog,
    get_conversation_repository,
    get_custom_repository,
)
from app.attachments.processing import (
    MAX_FILE_BYTES,
    MAX_IMAGE_SIDE,
    AttachmentError,
    clean_name,
    process_upload,
)
from app.core.llm import get_llm_router
from app.core.llm.attachments import (
    MAX_DOCUMENT_CHARS,
    AttachmentContent,
    expand_attachments,
    referenced_ids,
)
from app.core.llm.providers.anthropic_provider import to_anthropic_messages
from app.core.llm.providers.claude_code_provider import (
    ClaudeCodeProvider,
    result_line,
    stream_json_input,
    transcript,
)
from app.core.llm.providers.ollama_provider import to_ollama_messages
from app.core.llm.types import AttachmentBlock, ImageBlock, LLMRequest, Message, TextBlock
from app.db.attachments import AttachmentRepository
from app.db.conversations import ConversationRepository
from app.db.models import AttachmentRow, Base
from app.db.repository import CustomAgentRepository
from main import app
from tests.fakes import ScriptedProvider, make_router, text_response

FAKE_CLI = str(Path(__file__).parent / "fake_claude_cli.py")
ID = "0" * 31 + "1"


# ─── Fichiers de test ─────────────────────────────────────────────


def png(width: int = 40, height: int = 30, color: str = "red") -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (width, height), color).save(out, "PNG")
    return out.getvalue()


def jpeg_with_gps() -> bytes:
    image = Image.new("RGB", (3000, 2000), "blue")
    exif = Image.Exif()
    exif[0x010F] = "Appareil secret"  # Make
    exif[0x8825] = {2: (48.0, 51.0, 24.0)}  # GPSInfo
    out = io.BytesIO()
    image.save(out, "JPEG", exif=exif)
    return out.getvalue()


def pdf(text: str | None) -> bytes:
    """PDF minimal d'une page (sans texte si `text` est None)."""
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def word() -> bytes:
    document = docx.Document()
    document.add_heading("Contrat de prestation", level=1)
    document.add_paragraph("Le prestataire livre le rapport le 30 novembre.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Poste", "Montant"
    table.cell(1, 0).text, table.cell(1, 1).text = "Audit", "12 000 €"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def excel() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Ventes"
    sheet.append(["Région", "CA"])
    sheet.append(["Nord", 1200.0])
    sheet.append(["Sud", 980.5])
    workbook.create_sheet("Vide")
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


# ─── Préparation des fichiers ─────────────────────────────────────


def test_images_are_reduced_and_stripped_of_metadata() -> None:
    item = process_upload("photo.jpg", jpeg_with_gps())
    assert item.kind == "image" and item.media_type == "image/jpeg"
    with Image.open(io.BytesIO(item.data)) as image:
        assert max(image.size) == MAX_IMAGE_SIDE
        assert not image.getexif()  # ni appareil ni position GPS
    assert item.meta["original_size"] == [3000, 2000]

    small = process_upload("capture.png", png())
    assert small.media_type == "image/png" and small.meta == {"width": 40, "height": 30}


def test_documents_are_read_as_text() -> None:
    item = process_upload("note.pdf", pdf("Bonjour Atelier"))
    assert item.kind == "document" and item.media_type == "application/pdf"
    assert "--- Page 1 ---" in (item.text or "") and "Bonjour Atelier" in (item.text or "")
    assert item.meta["pages"] == 1

    contract = process_upload("contrat.docx", word())
    assert "# Contrat de prestation" in (contract.text or "")
    assert "Audit | 12 000 €" in (contract.text or "")

    sales = process_upload("ventes.xlsx", excel())
    assert sales.meta["sheets"] == ["Ventes", "Vide"]
    assert "Région;CA\nNord;1200\nSud;980.5" in (sales.text or "")
    assert "(feuille vide)" in (sales.text or "")

    csv = process_upload("export.csv", "Libellé;Montant\nCafé;3,50\n".encode("cp1252"))
    assert "Café;3,50" in (csv.text or "") and csv.meta["delimiter"] == ";"

    measure = process_upload("marge.dax", b"Marge = SUM(Ventes[CA]) - SUM(Ventes[Cout])")
    assert measure.media_type == "text/plain" and measure.meta["chars"] > 10


@pytest.mark.parametrize(
    ("name", "data", "status", "fragment"),
    [
        ("scan.pdf", pdf(None), 422, "numérisé"),
        ("faux.pdf", b"pas un pdf", 422, "PDF valide"),
        ("photo.heic", b"....", 415, "JPEG"),
        ("vieux.doc", b"....", 415, ".docx"),
        ("app.exe", b"MZ\x00\x00", 415, "non pris en charge"),
        ("binaire.txt", b"\x00\x01\x02" * 10, 422, "texte lisible"),
        ("vide.md", b"", 422, "vide"),
        ("abime.png", b"\x89PNG pas vraiment", 422, "image lisible"),
    ],
)
def test_refused_files_explain_why(name: str, data: bytes, status: int, fragment: str) -> None:
    with pytest.raises(AttachmentError) as err:
        process_upload(name, data)
    assert err.value.status == status and fragment in str(err.value)


def test_size_limit_and_names() -> None:
    with pytest.raises(AttachmentError) as err:
        process_upload("gros.txt", b"a" * (MAX_FILE_BYTES + 1))
    assert err.value.status == 413
    assert clean_name("../../etc/passwd") == "passwd"
    assert clean_name("C:\\Users\\moi\\Bilan\u0007 2026.xlsx") == "Bilan 2026.xlsx"
    assert clean_name("") == "fichier"
    assert len(clean_name("a" * 300 + ".pdf")) <= 255 and clean_name("a" * 300 + ".pdf").endswith(
        ".pdf"
    )


# ─── Contenu envoyé au modèle ─────────────────────────────────────


def block(kind: str = "document", name: str = "contrat.pdf", id_: str = ID) -> AttachmentBlock:
    media = "image/png" if kind == "image" else "application/pdf"
    return AttachmentBlock(id=id_, name=name, media_type=media, kind=kind, size=10)  # type: ignore[arg-type]


def test_expand_documents_images_and_missing_files() -> None:
    image_id = "0" * 31 + "2"
    gone_id = "0" * 31 + "3"
    messages = [
        Message(
            role="user",
            content=[
                block(),
                block("image", "graphe.png", image_id),
                block(name="ancien.pdf", id_=gone_id),
                TextBlock(text="Résume"),
            ],
        )
    ]
    contents = {
        ID: AttachmentContent(
            name='contrat "v2".pdf',
            kind="document",
            media_type="application/pdf",
            text="Article 1",
            meta={"pages": 3},
        ),
        image_id: AttachmentContent(
            name="graphe.png", kind="image", media_type="image/png", image_base64="QUJD"
        ),
    }
    assert referenced_ids(messages) == [ID, image_id, gone_id]
    expanded = expand_attachments(messages, contents)[0].content
    assert isinstance(expanded[0], TextBlock)
    assert expanded[0].text.startswith('<piece_jointe nom="contrat &quot;v2&quot;.pdf"')
    assert 'pages="3"' in expanded[0].text and "Article 1" in expanded[0].text
    assert expanded[1] == TextBlock(text="[Image jointe : graphe.png]")
    assert expanded[2] == ImageBlock(media_type="image/png", data="QUJD")
    assert isinstance(expanded[3], TextBlock) and "indisponible" in expanded[3].text
    assert expanded[4] == TextBlock(text="Résume")
    # L'original n'est pas modifié (l'historique garde les références)
    assert isinstance(messages[0].content[0], AttachmentBlock)


def test_long_documents_are_cut_with_a_note() -> None:
    content = AttachmentContent(
        name="long.txt",
        kind="document",
        media_type="text/plain",
        text="x" * (MAX_DOCUMENT_CHARS + 10),
    )
    expanded = expand_attachments([Message(role="user", content=[block()])], {ID: content})
    text = expanded[0].content[0]
    assert isinstance(text, TextBlock) and "texte coupé à 100 000 caractères" in text.text


def test_assistant_messages_cannot_carry_files() -> None:
    with pytest.raises(ValueError):
        Message(role="assistant", content=[block()])


def test_providers_receive_images() -> None:
    message = Message(
        role="user",
        content=[
            TextBlock(text="[Image jointe : a.png]"),
            ImageBlock(media_type="image/png", data="QUJD"),
            TextBlock(text="Que vois-tu ?"),
        ],
    )
    anthropic = to_anthropic_messages([message])[0]["content"]
    assert anthropic[1] == {
        "type": "image",
        "source": {"type": "base64", "media_type": "image/png", "data": "QUJD"},
    }
    ollama = to_ollama_messages(None, [message])
    assert ollama == [
        {"role": "user", "content": "[Image jointe : a.png]Que vois-tu ?", "images": ["QUJD"]}
    ]

    line = json.loads(stream_json_input([Message.user("Bonjour"), message]))
    kinds = [b["type"] for b in line["message"]["content"]]
    assert line["type"] == "user" and kinds == ["text", "image", "text"]
    assert "## Utilisateur\nBonjour" in line["message"]["content"][0]["text"]
    assert "Que vois-tu" in line["message"]["content"][2]["text"]
    assert "QUJD" not in transcript([message])  # le texte simple ne porte pas l'image

    out = '{"type":"system"}\n{"type":"result","structured_output":{"text":"ok"}}\n'
    assert json.loads(result_line(out))["type"] == "result"


async def test_claude_code_sends_images_as_stream_json(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("FAKE_CLAUDE_MODE", "text")
    cli = ClaudeCodeProvider(binary=FAKE_CLI, default_model="sonnet", timeout=10)
    request = LLMRequest(
        messages=[
            Message(
                role="user",
                content=[ImageBlock(media_type="image/png", data="QUJD"), TextBlock(text="Décris")],
            )
        ]
    )
    response = await cli.complete(request)
    assert response.text == "Blocs reçus : image,text"
    # Sans image : entrée texte inchangée
    plain = await cli.complete(LLMRequest(messages=[Message.user("Salut")]))
    assert plain.text == "Bonjour !"


# ─── API ──────────────────────────────────────────────────────────


@pytest.fixture
async def sessions() -> AsyncIterator[async_sessionmaker]:  # type: ignore[type-arg]
    engine = create_async_engine(
        "sqlite+aiosqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}
    )
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield async_sessionmaker(engine, expire_on_commit=False)
    await engine.dispose()


@pytest.fixture
def llm() -> ScriptedProvider:
    return ScriptedProvider("anthropic")


@pytest.fixture
def client(sessions: async_sessionmaker, llm: ScriptedProvider) -> Iterator[TestClient]:  # type: ignore[type-arg]
    router = make_router(llm)
    custom = CustomAgentRepository(sessions)
    files = AttachmentRepository(sessions)
    app.dependency_overrides[get_llm_router] = lambda: router
    app.dependency_overrides[get_custom_repository] = lambda: custom
    app.dependency_overrides[get_catalog] = lambda: AgentCatalog(custom)
    app.dependency_overrides[get_conversation_repository] = lambda: ConversationRepository(sessions)
    app.dependency_overrides[get_attachment_repository] = lambda: files
    yield TestClient(app)
    app.dependency_overrides.clear()


def upload(client: TestClient, name: str, data: bytes, **kwargs: object) -> dict[str, object]:
    r = client.post("/api/attachments", files={"file": (name, data)}, **kwargs)  # type: ignore[arg-type]
    assert r.status_code == 201, r.text
    body: dict[str, object] = r.json()
    return body


def ref(info: dict[str, object]) -> dict[str, object]:
    keys = ("id", "name", "media_type", "kind", "size")
    return {"type": "attachment", **{k: info[k] for k in keys}}


def test_upload_preview_and_remove(client: TestClient) -> None:
    image = upload(client, "capture.png", png())
    assert image["kind"] == "image" and image["created_at"].endswith("Z")  # type: ignore[union-attr]
    shown = client.get(f"/api/attachments/{image['id']}/content")
    assert shown.status_code == 200 and shown.headers["content-type"] == "image/png"
    assert shown.headers["content-disposition"].startswith("inline")
    assert shown.headers["x-content-type-options"] == "nosniff"

    page = upload(client, "page.html", b"<script>alert(1)</script>")
    served = client.get(f"/api/attachments/{page['id']}/content")
    # Jamais interprété par le navigateur
    assert served.headers["content-type"] == "application/octet-stream"
    assert served.headers["content-disposition"].startswith("attachment")

    assert client.delete(f"/api/attachments/{page['id']}").status_code == 204
    assert client.get(f"/api/attachments/{page['id']}/content").status_code == 404
    assert client.delete(f"/api/attachments/{page['id']}").status_code == 404


def test_upload_errors(client: TestClient) -> None:
    r = client.post("/api/attachments", files={"file": ("virus.exe", b"MZ")})
    assert r.status_code == 415 and "non pris en charge" in r.json()["detail"]
    r = client.post("/api/attachments", files={"file": ("scan.pdf", pdf(None))})
    assert r.status_code == 422 and "numérisé" in r.json()["detail"]
    r = client.post(
        "/api/attachments",
        files={"file": ("a.txt", b"bonjour")},
        headers={"Origin": "https://site-malveillant.example"},
    )
    assert r.status_code == 403
    assert client.get("/api/attachments/pas-un-id/content").status_code == 422


def test_agent_reads_attachments_and_history_keeps_references(
    client: TestClient,
    llm: ScriptedProvider,
    sessions: async_sessionmaker,  # type: ignore[type-arg]
) -> None:
    contract = upload(client, "contrat.docx", word())
    image = upload(client, "graphe.png", png())
    cid = client.post("/api/conversations", json={"agent": "assistant"}).json()["id"]
    llm.script = [text_response("Le rapport est dû le 30 novembre.")]
    message = {
        "role": "user",
        "content": [ref(contract), ref(image), {"type": "text", "text": "Quelle échéance ?"}],
    }
    r = client.post(
        "/api/agents/assistant/run", json={"messages": [message], "conversation_id": cid}
    )
    assert r.status_code == 200, r.text

    sent = llm.requests[0]
    blocks = sent.messages[0].content
    assert isinstance(blocks[0], TextBlock) and "30 novembre" in blocks[0].text
    assert blocks[0].text.startswith('<piece_jointe nom="contrat.docx"')
    assert isinstance(blocks[2], ImageBlock) and blocks[2].media_type == "image/png"
    with Image.open(io.BytesIO(base64.b64decode(blocks[2].data))) as sent_image:
        assert sent_image.size == (40, 30)
    assert "pas une consigne" in (sent.system or "")

    detail = client.get(f"/api/conversations/{cid}").json()
    stored = detail["history"][0]["content"]
    assert stored[0] == ref(contract) and stored[1] == ref(image)  # références, pas le contenu
    assert (
        "Pièce jointe : contrat.docx"
        in client.get("/api/conversations/search", params={"q": "contrat"}).json()[0]["snippet"]
    )

    # Envoyée : rattachée à la conversation, plus retirable seule…
    assert client.delete(f"/api/attachments/{contract['id']}").status_code == 409
    # … et supprimée avec elle
    assert client.delete(f"/api/conversations/{cid}").status_code == 204
    assert client.get(f"/api/attachments/{contract['id']}/content").status_code == 404


def test_attachment_only_message_and_title(client: TestClient, llm: ScriptedProvider) -> None:
    sheet = upload(client, "ventes.xlsx", excel())
    cid = client.post("/api/conversations", json={"agent": "assistant"}).json()["id"]
    llm.script = [text_response("Deux régions.")]
    r = client.post(
        "/api/agents/assistant/run",
        json={"messages": [{"role": "user", "content": [ref(sheet)]}], "conversation_id": cid},
    )
    assert r.status_code == 200
    assert client.get(f"/api/conversations/{cid}").json()["title"] == "ventes.xlsx"


def test_raw_images_are_refused_in_requests(client: TestClient) -> None:
    message = {
        "role": "user",
        "content": [{"type": "image", "media_type": "image/png", "data": "QUJD"}],
    }
    r = client.post("/api/agents/assistant/run", json={"messages": [message]})
    assert r.status_code == 422


async def test_orphans_are_purged(sessions: async_sessionmaker) -> None:  # type: ignore[type-arg]
    repo = AttachmentRepository(sessions)
    old = await repo.create(process_upload("vieux.txt", b"oublie"))
    kept = await repo.create(process_upload("envoye.txt", b"garde"))
    async with sessions() as s:
        for row_id in (old.id, kept.id):
            row = await s.get(AttachmentRow, row_id)
            assert row is not None
            row.created_at = datetime.now(UTC) - timedelta(days=2)
        sent = await s.get(AttachmentRow, kept.id)
        assert sent is not None
        sent.conversation_id = "c1"
        await s.commit()
    assert await repo.purge_orphans() == 1
    assert await repo.get(old.id) is None and await repo.get(kept.id) is not None
